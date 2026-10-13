# -*- coding: utf-8 -*-
"""Genera PITCH_DECK_SGR_CONV56_AIRE.pptx (15 diapositivas) con la paleta del proyecto, desde gen_x03.SLIDES."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gen_x03
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TINTA = RGBColor(0x2A, 0x34, 0x39); AZUL = RGBColor(0x66, 0x99, 0xCC); MEDIO = RGBColor(0x9E, 0xB9, 0xD4); CLARO = RGBColor(0xB9, 0xD9, 0xEB); GRIS = RGBColor(0xF5, 0xF5, 0xF5)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
def caja(sl, x, y, w, h, fill=None, line=None):
    s = sl.shapes.add_shape(1, x, y, w, h)
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    return s
def texto(sl, x, y, w, h, txt, size=18, bold=False, color=TINTA, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "DejaVu Sans"
    return tf
for n, titulo, items, nota in gen_x03.SLIDES:
    sl = prs.slides.add_slide(blank)
    caja(sl, 0, 0, prs.slide_width, Inches(1.15), fill=CLARO)
    caja(sl, 0, Inches(1.15), prs.slide_width, Emu(30000), fill=AZUL)
    texto(sl, Inches(0.5), Inches(0.18), Inches(11.6), Inches(0.8), titulo, 30, True, anchor=MSO_ANCHOR.MIDDLE)
    texto(sl, Inches(12.0), Inches(0.3), Inches(1.0), Inches(0.5), f"{n}/15", 14, False, align=PP_ALIGN.RIGHT)
    tb = sl.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.2)); tf = tb.text_frame; tf.word_wrap = True
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if k == 0 else 14)
        r = p.add_run(); r.text = (it if n == 1 else "•  " + it); r.font.size = Pt(24 if n > 1 else 22); r.font.bold = (n == 1 and k == 0)
        r.font.color.rgb = TINTA; r.font.name = "DejaVu Sans"
    sl.notes_slide.notes_text_frame.text = nota
    texto(sl, Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.35), "SGR_CONV56_AIRE · Borrador con vacíos · 13 de octubre de 2026", 10, False, TINTA)
out = os.path.join(RAIZ, "out", "PITCH_DECK_SGR_CONV56_AIRE.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out); print("pptx", out, len(prs.slides))
