# -*- coding: utf-8 -*-
"""Posprocesado Word del documento único (estilo APA 7, plantilla visual del usuario)."""
import docx, re, sys, os, json
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from xml.sax.saxutils import escape

TINTA = "2A3439"; AZUL = "6699CC"; MEDIO = "9EB9D4"; CLARO = "B9D9EB"; CAJA = "F7F7F7"; GRIS = "F5F5F5"
ORD = {
 'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription',
 'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark',
 'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
 'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath',
 'style': 'name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr',
 'trPr': 'cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden',
 'settings': 'writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids mathPr attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags decimalSymbol listSeparator',
}
def local(tag): return tag.split('}')[1]
def reorder(root):
    for name, seq in ORD.items():
        idx = {n: i for i, n in enumerate(seq.split())}
        for el in root.iter(qn('w:' + name)):
            kids = list(el)
            keyed = sorted(kids, key=lambda k: idx.get(local(k.tag), len(idx)) if isinstance(k.tag, str) else len(idx))
            if keyed != kids:
                for k in kids: el.remove(k)
                for k in keyed: el.append(k)

def ncols(t):
    return len(t.find(qn('w:tblGrid')).findall(qn('w:gridCol')))

def shade(tc, fill):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None: tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)

def set_cell_border(tc, **kw):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None: tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
    b = tcPr.find(qn('w:tcBorders'))
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge, (sz, color) in kw.items():
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        for old in b.findall(qn('w:' + edge)): b.remove(old)
        b.append(e)

def tabla_apa(tbl, d, styles_id):
    n = ncols(tbl)
    tblPr = tbl.find(qn('w:tblPr'))
    for tag in ('w:tblBorders', 'w:tblCellMar', 'w:tblLayout'):
        for old in tblPr.findall(qn(tag)): tblPr.remove(old)
    # APA: sin líneas verticales; líneas horizontales solo arriba, bajo encabezados y al cierre
    tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="0" w:color="{TINTA}"/>'
                           f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="{TINTA}"/></w:tblBorders>'))
    tblPr.append(parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="30" w:type="dxa"/><w:left w:w="70" w:type="dxa"/><w:bottom w:w="30" w:type="dxa"/><w:right w:w="70" w:type="dxa"/></w:tblCellMar>'))
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    st = 'Tabla9' if n <= 6 else ('Tabla8' if n <= 10 else ('Tabla7' if n <= 14 else 'Tabla65'))
    if re.search(r'[0-9a-f]{64}', ''.join(x.text or '' for x in tbl.iter(qn('w:t')))) and n < 9: st = 'Tabla8'
    rows = tbl.findall(qn('w:tr'))
    for ri, tr in enumerate(rows):
        trPr = tr.find(qn('w:trPr'))
        if trPr is None: trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
        trPr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trPr.append(OxmlElement('w:tblHeader'))
        for tc in tr.findall(qn('w:tc')):
            if ri == 0:
                shade(tc, CLARO); set_cell_border(tc, bottom=(12, AZUL))
            elif ri % 2 == 0:
                shade(tc, CAJA)
            for p in tc.findall(qn('w:p')):
                pPr = p.find(qn('w:pPr'))
                if pPr is None: pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
                ps = pPr.find(qn('w:pStyle'))
                if ps is None: ps = OxmlElement('w:pStyle'); pPr.insert(0, ps)
                ps.set(qn('w:val'), styles_id[st])
                if ri == 0:
                    for r in p.findall(qn('w:r')):
                        rPr = r.find(qn('w:rPr'))
                        if rPr is None: rPr = OxmlElement('w:rPr'); r.insert(0, rPr)
                        if rPr.find(qn('w:b')) is None: rPr.append(OxmlElement('w:b'))

def field(p, instr, texto='1', size=10):
    for kind, txt in (('begin', None), ('instr', instr), ('separate', None), ('text', texto), ('end', None)):
        r = p.add_run(); r.font.size = Pt(size)
        if kind == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {txt} '; r._r.append(it)
        elif kind == 'text': r.text = txt
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)

SECT_P = '<w:p {ns}><w:pPr><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>{sect}</w:pPr></w:p>'
def sect_xml(land):
    if land:
        return ('<w:sectPr><w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>'
                '<w:pgMar w:top="1440" w:right="1000" w:bottom="1440" w:left="1000" w:header="680" w:footer="680" w:gutter="0"/></w:sectPr>')
    return ('<w:sectPr><w:pgSz w:w="12240" w:h="15840"/>'
            '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="680" w:footer="680" w:gutter="0"/></w:sectPr>')
def mk_sect_par(land):
    return parse_xml(SECT_P.format(ns=nsdecls('w'), sect=sect_xml(land)))

def postprocesar(src, dst, portada=None, wide_cols=12):
    d = docx.Document(src); body = d.element.body
    styles_id = {s.name: s.style_id for s in d.styles}
    # listas más compactas
    for lvl in d.part.numbering_part.element.iter(qn('w:lvl')):
        ppr = lvl.find(qn('w:pPr'))
        if ppr is None: continue
        ind = ppr.find(qn('w:ind'))
        if ind is None: continue
        il = int(lvl.get(qn('w:ilvl')))
        ind.set(qn('w:left'), str(567 + 340 * il)); ind.set(qn('w:hanging'), '283')
        if ind.get(qn('w:firstLine')) is not None: del ind.attrib[qn('w:firstLine')]
    # tablas APA
    for t in body.findall(qn('w:tbl')): tabla_apa(t, d, styles_id)
    # apaisado para tablas muy anchas
    children = list(body)
    wide = [i for i, el in enumerate(children) if el.tag == qn('w:tbl') and ncols(el) >= wide_cols]
    for i in sorted(wide, reverse=True):
        el = children[i]
        # incluir rótulo y título previos
        first = el
        j = i - 1
        while j >= 0 and children[j].tag == qn('w:p') and (children[j].find(qn('w:pPr')) is not None) and \
                children[j].find(qn('w:pPr')).find(qn('w:pStyle')) is not None and \
                children[j].find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val')) in (styles_id['TablaRotulo'], styles_id['TablaTitulo']):
            first = children[j]; j -= 1
        last = el
        nx = el.getnext()
        if nx is not None and nx.tag == qn('w:p') and nx.find(qn('w:pPr')) is not None and nx.find(qn('w:pPr')).find(qn('w:pStyle')) is not None \
                and nx.find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val')) == styles_id['TablaNota']:
            last = nx
        first.addprevious(mk_sect_par(False)); last.addnext(mk_sect_par(True))
    # sectPr final vertical
    fin = body.find(qn('w:sectPr'))
    for ch in list(fin):
        if ch.tag in (qn('w:pgSz'), qn('w:pgMar')): fin.remove(ch)
    fin.insert(0, parse_xml(f'<w:pgSz {nsdecls("w")} w:w="12240" w:h="15840"/>'))
    fin.insert(1, parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="680" w:footer="680" w:gutter="0"/>'))
    # encabezado con número de página en la esquina superior derecha (desde la portada)
    sec = d.sections[0]
    hp = sec.header.paragraphs[0]; hp.text = ''; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    field(hp, 'PAGE', '1', 11)
    sec.footer.paragraphs[0].text = ''
    # portada
    if portada:
        first = None
        sdt = body.find(qn('w:sdt'))
        anchor = sdt if sdt is not None else body[0]
        for kind, txt, kw in portada:
            if kind == 'p':
                p = docx.text.paragraph.Paragraph(OxmlElement('w:p'), d._body)
                anchor.addprevious(p._p)
                r = p.add_run(txt); r.font.size = Pt(kw.get('size', 12)); r.bold = kw.get('bold', False); r.italic = kw.get('italic', False)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format; pf.space_before = Pt(kw.get('before', 0)); pf.space_after = Pt(kw.get('after', 6)); pf.line_spacing = 1.15
                pf.first_line_indent = Cm(0)
            elif kind == 'break':
                bp = parse_xml(f'<w:p {nsdecls("w")}><w:r><w:br w:type="page"/></w:r></w:p>'); anchor.addprevious(bp)
    # limpieza de fldChar dirty
    for fc in body.iter(qn('w:fldChar')):
        if fc.get(qn('w:dirty')) is not None: del fc.attrib[qn('w:dirty')]
    # orden XML
    reorder(d.element); reorder(d.settings.element); reorder(d.styles.element)
    for rp in d.styles.element.iter(qn('w:rPr')):
        for k in rp:
            if k.tail and k.tail.strip(): k.tail = None
    for ns_ in d.part.numbering_part.element.iter(qn('w:nsid'), qn('w:tmpl')):
        v = ns_.get(qn('w:val'))
        if len(v) != 8: ns_.set(qn('w:val'), v.upper().zfill(8))
    cp = d.core_properties
    cp.title = 'Proyecto completo SGR_CONV56_AIRE'; cp.subject = 'GENESIS-CTeI 36.10 RC2'; cp.author = ''; cp.language = 'es-CO'
    d.save(dst)
    return dst

# ---------------------------------------------------------------- índices (tabla de contenido, de tablas y de figuras)
def titulos(d):
    body = d.element.body; ids = {s.style_id: s.name for s in d.styles}
    out = []
    for p in body.findall(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        if pPr is None: continue
        ps = pPr.find(qn('w:pStyle'))
        if ps is None: continue
        nm = ids.get(ps.get(qn('w:val')), '')
        m = re.fullmatch(r'Heading ([123])', nm)
        if not m: continue
        txt = ''.join(t.text or '' for t in p.iter(qn('w:t')))
        pv = p.getprevious(); bm = pv.get(qn('w:name')) if pv is not None and pv.tag == qn('w:bookmarkStart') else None
        out.append((int(m.group(1)), txt, bm))
    return out

def rotulos(d, estilo):
    """Devuelve [(rotulo 'Tabla 8.3', titulo, parrafo_rotulo)] para el estilo TablaRotulo o FiguraRotulo."""
    body = d.element.body; ids = {s.style_id: s.name for s in d.styles}
    res = []
    ps_ = list(body.findall(qn('w:p')))
    for k, p in enumerate(ps_):
        pPr = p.find(qn('w:pPr'))
        if pPr is None: continue
        ps = pPr.find(qn('w:pStyle'))
        if ps is None or ids.get(ps.get(qn('w:val'))) != estilo: continue
        rot = ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip()
        # el título es el siguiente párrafo
        nx = p.getnext(); tit = ''
        if nx is not None and nx.tag == qn('w:p'): tit = ''.join(t.text or '' for t in nx.iter(qn('w:t'))).strip()
        res.append((rot, tit, p))
    return res

def marcar_bookmarks(d):
    """Añade marcadores a los rótulos de tablas y figuras (para enlaces y PAGEREF)."""
    n = 9000
    for estilo in ('TablaRotulo', 'FiguraRotulo'):
        for rot, tit, p in rotulos(d, estilo):
            name = 'rot_' + re.sub(r'[^A-Za-z0-9]', '_', rot)
            n += 1
            s = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{n}" w:name="{name}"/>'); e = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{n}"/>')
            ppr = p.find(qn('w:pPr'))
            if ppr is not None: ppr.addnext(s)
            else: p.insert(0, s)
            p.append(e)

def entradas_indice(d, estilo):
    return [(rot, tit, 'rot_' + re.sub(r'[^A-Za-z0-9]', '_', rot)) for rot, tit, p in rotulos(d, estilo)]

def construir_toc(d, heads, pages, tabpos=9360):
    body = d.element.body; ids = {s.name: s.style_id for s in d.styles}
    sdt = body.find(qn('w:sdt')); content = sdt.find(qn('w:sdtContent'))
    ps = content.findall(qn('w:p'))
    if len(ps) > 1:
        for q in ps[1:]: content.remove(q)
    W = nsdecls('w')
    for k, ((lvl, txt, bm), pg) in enumerate(zip(heads, pages)):
        sid = ids[f'TOC {lvl}']
        pre = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>') if k == 0 else ''
        lo = f'<w:hyperlink w:anchor="{escape(bm)}" w:history="1">' if bm else ''
        lc = '</w:hyperlink>' if bm else ''
        x = (f'<w:p {W}><w:pPr><w:pStyle w:val="{sid}"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{tabpos}"/></w:tabs></w:pPr>{pre}{lo}'
             f'<w:r><w:t xml:space="preserve">{escape(txt)}</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>{pg}</w:t></w:r>{lc}</w:p>')
        content.append(parse_xml(x))
    content.append(parse_xml(f'<w:p {W}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'))

def construir_lista(d, marcador, entradas, paginas, tabpos=9360):
    """Sustituye el párrafo marcador (@@INDICE_TABLAS@@ o @@INDICE_FIGURAS@@) por entradas con enlace y PAGEREF."""
    body = d.element.body; ids = {s.name: s.style_id for s in d.styles}
    tgt = None
    for p in body.findall(qn('w:p')):
        if ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip() == marcador: tgt = p; break
    if tgt is None: return False
    W = nsdecls('w'); sid = ids['Indice']
    for (rot, tit, bm), pg in zip(entradas, paginas):
        txt = f'{rot}. {tit}'
        x = (f'<w:p {W}><w:pPr><w:pStyle w:val="{sid}"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{tabpos}"/></w:tabs>'
             f'<w:spacing w:before="40" w:after="40"/><w:ind w:left="1134" w:hanging="1134"/></w:pPr>'
             f'<w:hyperlink w:anchor="{escape(bm)}" w:history="1"><w:r><w:t xml:space="preserve">{escape(txt)}</w:t></w:r><w:r><w:tab/></w:r>'
             f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGEREF {escape(bm)} \\h </w:instrText></w:r>'
             f'<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>{pg}</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>')
        tgt.addprevious(parse_xml(x))
    body.remove(tgt)
    return True
