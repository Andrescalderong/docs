# -*- coding: utf-8 -*-
"""Libro de respaldo del Anexo A: variables, 210 fichas de dirección, matrices de los cinco escenarios (con totales por fórmula),
clases, potencias de S2 y ciclos. Uso: python3 build/gen_xlsx_anexoA.py SALIDA.xlsx"""
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__)); A2 = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(A2, "v2"))
import causal as C
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

TINTA = "2A3439"; CLARO = "B9D9EB"; MEDIO = "9EB9D4"; AZUL = "6699CC"
HF = Font(name="Arial", bold=True, color=TINTA); BF = Font(name="Arial", color=TINTA, size=10)
HFILL = PatternFill("solid", fgColor=CLARO)
FILL = {1: PatternFill("solid", fgColor=CLARO), 2: PatternFill("solid", fgColor=MEDIO), 3: PatternFill("solid", fgColor=AZUL)}
THIN = Side(style="thin", color=MEDIO); BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NAT = {"DOCUMENTADO": "Con respaldo documental", "INFERIDO": "Inferido de un mecanismo", "PROPUESTO": "Propuesta del equipo formulador", "SUPUESTO": "Supuesto de diseño"}

def cab(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t); c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDE
    if anchos:
        for j, w in enumerate(anchos, 1): ws.column_dimensions[L(j)].width = w

def main(out):
    wb = Workbook()
    ws = wb.active; ws.title = "Léame"
    filas = [
        "Anexo A del Anexo 2. Libro de respaldo del análisis estructural",
        "",
        "Contenido: variables, 210 direcciones con su ficha, matrices de los escenarios S0 a S4, clases, potencias de S2 y ciclos de S2.",
        "Los valores de influencia son juicios de formulación en escala ordinal de 0 a 3. Ningún panel externo los ha adjudicado.",
        "Las relaciones son hipótesis de influencia estructural y no efectos estimados.",
        "Los totales por fila (TA) y por columna (TP) de las matrices se calculan con fórmulas.",
        "S2 es la matriz de trabajo. S3 y S4 son envolventes de sensibilidad (mínimo y máximo admisible de cada vínculo).",
    ]
    for i, t in enumerate(filas, 1):
        c = ws.cell(i, 1, t); c.font = Font(name="Arial", bold=(i == 1), size=12 if i == 1 else 10, color=TINTA)
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Variables")
    cab(ws, 1, ["ID", "Variable", "Definición operacional (unidad)", "Bloque"], [8, 34, 70, 12])
    for i, v in enumerate(C.R["variables"], 2):
        for j, x in enumerate([v["id"], C.NOM[v["id"]], v["definicion"].replace("PM2.5", "PM2,5"), C.BLOQUE[v["bloque"]]], 1):
            c = ws.cell(i, j, x); c.font = BF; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BORDE

    ws = wb.create_sheet("Fichas_210")
    cab(ws, 1, ["Origen", "Destino", "Peso", "Dominio admisible", "Signo", "Confianza", "Naturaleza del juicio", "Mecanismo o razón del cero", "Explicación rival", "Respaldo"],
        [8, 8, 7, 12, 10, 11, 28, 70, 60, 36])
    r = 2
    for fi in C.F:
        for d in fi["direcciones"]:
            va = d["valores_admisibles"]; dom = str(va[0]) if len(va) == 1 else f"{min(va)} a {max(va)}"
            fila = [d["origen"], d["destino"], d["valor_propuesto"], dom, C.SIGNO.get(d["signo"], "No aplica" if d["signo"] == "NO_APLICA" else d["signo"]),
                    C.CONF[d["confianza"]["nivel"]], NAT.get(d["naturaleza"], d["naturaleza"]), C._limpiar(d["mecanismo_o_razon_cero"]),
                    C._limpiar(d["explicacion_rival"]), C._respaldo(d)]
            for j, x in enumerate(fila, 1):
                c = ws.cell(r, j, x); c.font = BF; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BORDE
            r += 1
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:J{r - 1}"

    for k, nm in C.ESC:
        ws = wb.create_sheet(f"Matriz_{k}")
        A = C.mat(k)
        cab(ws, 1, ["Origen"] + C.IDS + ["TA"], [9] + [6] * 15 + [7])
        for i in range(15):
            ws.cell(i + 2, 1, C.IDS[i]).font = HF
            for j in range(15):
                v = int(A[i, j]); c = ws.cell(i + 2, j + 2, v); c.font = BF; c.alignment = Alignment(horizontal="center"); c.border = BORDE
                if v: c.fill = FILL[v]
            c = ws.cell(i + 2, 17, f"=SUM(B{i + 2}:P{i + 2})"); c.font = HF; c.alignment = Alignment(horizontal="center")
        ws.cell(17, 1, "TP").font = HF
        for j in range(15):
            c = ws.cell(17, j + 2, f"=SUM({L(j + 2)}2:{L(j + 2)}16)"); c.font = HF; c.alignment = Alignment(horizontal="center")
        c = ws.cell(17, 17, "=SUM(Q2:Q16)"); c.font = HF; c.alignment = Alignment(horizontal="center")
        ws.cell(19, 1, C.ORIGEN[k]).font = Font(name="Arial", italic=True, size=10, color=TINTA)

    ws = wb.create_sheet("Clases")
    cab(ws, 1, ["Escenario", "ID", "Variable", "Influencia", "Dependencia", "Rango de influencia", "Rango de dependencia", "Clase por mediana", "Clase por media",
                "Rango de influencia indirecta", "Clase indirecta"], [10, 7, 32, 11, 12, 12, 12, 18, 18, 14, 18])
    r = 2
    for k, nm in C.ESC:
        s = C.S(k)
        for i, d in enumerate(s["directa"]):
            m = s["micmac"]["contraste"][i]
            fila = [k, d["id"], C.NOM[d["id"]], d["influencia"], d["dependencia"], d["rango_influencia"], d["rango_dependencia"], C.clase(d["clase_mediana"]),
                    C.clase(d["clase_media"]), m["rango_I_indirecto"], C.clase(m["clase_indirecta"])]
            for j, x in enumerate(fila, 1):
                c = ws.cell(r, j, x); c.font = BF; c.border = BORDE
            r += 1
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:K{r - 1}"

    ws = wb.create_sheet("Potencias_S2")
    pw = C.S("S2")["micmac"]["potencias"]
    cab(ws, 1, ["ID", "Variable"] + [f"I (A^{p['k']})" for p in pw], [7, 32] + [16] * len(pw))
    import numpy as np
    for i, v in enumerate(C.IDS):
        ws.cell(i + 2, 1, v).font = BF; ws.cell(i + 2, 2, C.NOM[v]).font = BF
        for j, p in enumerate(pw):
            c = ws.cell(i + 2, 3 + j, int(np.array(p["matriz"]).sum(1)[i])); c.font = BF; c.number_format = "#,##0"
    r0 = 20
    cab(ws, r0, ["ID", "Variable"] + [f"D (A^{p['k']})" for p in pw])
    for i, v in enumerate(C.IDS):
        ws.cell(r0 + 1 + i, 1, v).font = BF; ws.cell(r0 + 1 + i, 2, C.NOM[v]).font = BF
        for j, p in enumerate(pw):
            c = ws.cell(r0 + 1 + i, 3 + j, int(np.array(p["matriz"]).sum(0)[i])); c.font = BF; c.number_format = "#,##0"

    ws = wb.create_sheet("Ciclos_S2")
    cab(ws, 1, ["N.º", "Longitud", "Ciclo"], [7, 10, 60])
    for i, cy in enumerate(sorted(C.CICLOS_S2, key=lambda x: (len(x), x)), 2):
        ws.cell(i, 1, i - 1).font = BF; ws.cell(i, 2, len(cy)).font = BF
        ws.cell(i, 3, " → ".join(C.IDS[j] for j in cy) + " → " + C.IDS[cy[0]]).font = BF
    wb.save(out); print("ok", out)

if __name__ == "__main__":
    main(sys.argv[1])
