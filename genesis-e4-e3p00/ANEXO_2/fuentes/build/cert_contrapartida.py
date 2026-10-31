# -*- coding: utf-8 -*-
"""Certificaciones de contrapartida por integrante, según el Anexo 3 v7.
Uso: python3 build/cert_contrapartida.py carpeta_salida"""
import os, sys, docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
OUT = sys.argv[1]
TIT = "«Implementación de una red IoT híbrida Edge-Cloud y gemelo digital para la gestión anticipada del riesgo por material particulado posterremoto» (SIGP 143329)"
def doc(ent, carpeta, filas, notas):
    d = docx.Document(); s = d.styles["Normal"]; s.font.name = "Arial"; s.font.size = Pt(10.5)
    for sec in d.sections: sec.left_margin = sec.right_margin = Cm(2.5)
    def p(t, b=False, al=None, sz=10.5):
        q = d.add_paragraph(); r = q.add_run(t); r.bold = b; r.font.size = Pt(sz)
        if al is not None: q.alignment = al
    p(f"[Diligenciar: membrete de {ent}]", al=WD_ALIGN_PARAGRAPH.CENTER, sz=9)
    p("CERTIFICACIÓN DE CONTRAPARTIDA", True, WD_ALIGN_PARAGRAPH.CENTER, 12)
    p(f"El suscrito representante legal de {ent} [Diligenciar: NIT], en el marco de la Convocatoria 56 del Sistema General de Regalías (Asignación para la CTeI) y del proyecto {TIT}, cuya entidad proponente y ejecutora es la Universidad Santo Tomás, sede Villavicencio,")
    p("CERTIFICA:", True)
    p(f"Que {ent} se compromete a aportar al proyecto la contrapartida siguiente, cuantificada en términos monetarios, con la actividad que financia y el momento del aporte, según el Anexo 3 versión 7:")
    t = d.add_table(rows=1, cols=6); t.style = "Table Grid"
    for i, x in enumerate(["Código Anexo 3", "Concepto", "Modalidad", "Valor (COP)", "Actividad", "Momento"]): t.rows[0].cells[i].text = x
    for f in filas:
        c = t.add_row().cells
        for i, x in enumerate(f): c[i].text = x
    tot = sum(int(f[3].replace(".", "")) for f in filas)
    c = t.add_row().cells; c[1].text = "Total"; c[3].text = f"{tot:,}".replace(",", ".")
    for row in t.rows:
        for cell in row.cells:
            for r in cell.paragraphs[0].runs: r.font.size = Pt(9)
    d.add_paragraph()
    for n in notas: p(n)
    p("Que la presente certificación no constituye compromiso de posterior suscripción y se expide para ser cargada con la postulación.")
    p("Se expide en [Diligenciar: ciudad], el [Diligenciar: día] de noviembre de 2026.")
    p("\n\n______________________________"); p("[Diligenciar: nombre del representante legal]", True); p("Representante legal · C. C. [Diligenciar: número]"); p(f"{ent} · NIT [Diligenciar]")
    f = os.path.join(OUT, carpeta, f"Certificado_contrapartida_{ent.replace(' ', '_').replace('á','a')}.docx"); os.makedirs(os.path.dirname(f), exist_ok=True); d.save(f); print("ok", f)
VAL = "Que el valor se calcula por uso, con el costo de licencia o la depreciación proporcional al tiempo y a la capacidad asignados al proyecto, según el inventario y la memoria de valoración que se adjuntan [Diligenciar: inventario y memoria de valoración]. El aporte no incluye horas de funcionarios ni desembolsos en efectivo, y los bienes siguen siendo propiedad del municipio, sin transferencia."
for mun, eq, car in (("Municipio de Manizales", "EQ-24", "02_FIRMAR/B_Municipio_de_Manizales"), ("Municipio de Calarcá", "EQ-25", "02_FIRMAR/C_Municipio_de_Calarca")):
    doc(mun, car, [(eq, "Licencias del sistema de información geográfica corporativo para dos puestos de trabajo", "Especie (uso)", "20.000.000", "3.1 a 3.4", "Meses 3 a 24"),
                   (eq, "Alojamiento del tablero y de los datos del índice en los servidores y el almacenamiento del centro de datos municipal", "Especie (uso)", "15.000.000", "3.1 a 3.4", "Meses 3 a 24"),
                   (eq, "Equipos de cómputo y pantalla de la sala de gestión del riesgo para el uso institucional del tablero", "Especie (uso)", "15.000.000", "3.1 a 3.4", "Meses 3 a 24"),
                   (eq, "Conectividad, puntos de red y energía en las sedes municipales donde se instalen nodos", "Especie (uso)", "10.000.000", "3.1 a 3.4", "Meses 3 a 24")], [VAL])
doc("Avanciencia", "02_FIRMAR/D_Avanciencia", [("TH-22", "Horas de investigadores y del equipo de transferencia (vigilancia tecnológica, actualizaciones de libertad de operación y acompañamiento de convenios)", "Especie", "150.000.000", "1.1 (y 1.2, 1.4, 1.6, 3.3, 3.4)", "Meses 2 a 24")],
    ["Que el valor en especie se calcula con un método verificable [Diligenciar: costo institucional, tarifa o nómina certificada]."])
doc("Gestión y Medio Ambiente", "02_FIRMAR/E_Gestion_y_Medio_Ambiente", [("EQ-13", "Equipos de monitoreo de calidad del aire y apoyo técnico para el diagnóstico, la colocalización y la verificación en Manizales", "Especie", "110.000.000", "1.1 (y 1.3, 1.6)", "Meses 2 a 23"),
    ("VI-05", "Logística y viáticos de la instalación de los tramos 2 y 3", "Efectivo", "150.000.000", "2.3 (y 2.4)", "Meses 12 a 19")],
    ["Que el valor en especie se calcula con un método verificable [Diligenciar: costo institucional, tarifa o depreciación] y que los bienes se ubican en Manizales, lugar de ejecución.",
     "Que el aporte en efectivo se consignará en la cuenta que la entidad ejecutora disponga para el proyecto, según el cronograma financiero del Anexo 3, y que la entidad cuenta con los recursos para hacerlo [Diligenciar: soporte de disponibilidad]."])
