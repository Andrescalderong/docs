# -*- coding: utf-8 -*-
"""Resumen ejecutivo del proyecto SIGP 143329 para el alcalde de cada municipio (Anexo 3 v7).
Uso: python3 build/resumen_alcalde.py carpeta_salida"""
import os, sys
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = sys.argv[1]
GRAF = "/tmp/claude-0/-home-user-docs/8ca5c040-2adb-56b3-8e64-05d4e7e569ed/scratchpad/grafico_integrado.png"
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
TIT = ("Implementación de una red IoT híbrida Edge-Cloud y gemelo digital para la gestión anticipada del riesgo "
       "por material particulado posterremoto")

def base():
    d = docx.Document()
    s = d.styles["Normal"]; s.font.name = "Arial"; s.font.size = Pt(10)
    s.paragraph_format.space_after = Pt(4)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(1.8)
    return d

def p(d, t, b=False, s=10, al=None, after=4, color=None, i=False):
    q = d.add_paragraph(); r = q.add_run(t); r.bold = b; r.italic = i; r.font.size = Pt(s)
    if color: r.font.color.rgb = color
    q.paragraph_format.space_after = Pt(after)
    if al is not None: q.alignment = al
    return q

def h(d, t):
    q = p(d, t, b=True, s=12, color=AZUL, after=3); q.paragraph_format.space_before = Pt(8); return q

def bullet(d, t):
    q = d.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(2)
    r = q.add_run(t); r.font.size = Pt(10)

def sombra(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)

def tabla(d, enc, filas, anchos=None, fs=8.5):
    t = d.add_table(rows=1, cols=len(enc)); t.style = "Table Grid"
    for i, x in enumerate(enc):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(x); r.bold = True; r.font.size = Pt(fs); sombra(c, "DCE6F0")
    for f in filas:
        cs = t.add_row().cells
        for i, x in enumerate(f):
            cs[i].text = ""; r = cs[i].paragraphs[0].add_run(x); r.font.size = Pt(fs)
            if x.startswith("Total"): r.bold = True
    if anchos:
        for row in t.rows:
            for i, w in enumerate(anchos): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def resumen(mun, dep, eq, otro):
    d = base()
    p(d, "RESUMEN EJECUTIVO PARA EL DESPACHO DEL ALCALDE", b=True, s=9, color=AZUL)
    p(d, TIT, b=True, s=14, after=2)
    p(d, f"Convocatoria 56 del Sistema General de Regalías (Asignación para la CTeI) · Código SIGP 143329 · Municipio de {mun} ({dep}) · "
         "Cierre de la convocatoria: martes 3 de noviembre de 2026, 4:00 p. m.", s=9, i=True, after=8)

    h(d, "1. En una frase")
    p(d, f"El proyecto instala en {mun} una red de 50 sensores de material particulado (PM2,5 y PM10) en los sitios de demolición y remoción de escombros "
         "del sismo del 10 de agosto de 2026, verifica su error frente a equipos de referencia y entrega al consejo municipal de gestión del riesgo un "
         "tablero y un protocolo de alerta para proteger a trabajadores y vecinos del polvo, más un laboratorio que estima la exposición antes de cada demolición.")

    h(d, "2. Objetivo, alcance y justificación")
    p(d, "Objetivo general: reducir el error de las mediciones de PM2,5 y PM10 en los sitios de remoción de escombros y lograr que los consejos de gestión "
         "del riesgo las usen para decidir alertas, en 24 meses.", b=False)
    tabla(d, ["Objetivo específico", "Qué entrega", "Valor SGR (COP)"], [
        ("1. Corregir la medición y validar el modelo predictivo", "Colocalización con equipos de referencia, modelo de corrección validado en cada municipio, gemelo predictivo probado en al menos seis demoliciones reales", "7.807.068.548"),
        ("2. Red disponible y laboratorio de modelación", "100 nodos y 10 gateways con energía solar y respaldo satelital (50 y 5 por municipio), plataforma y laboratorio de modelación", "4.620.145.795"),
        ("3. Uso institucional y custodia", "Umbrales y protocolo de alerta co-diseñados con los consejos, 20 talleres, 9 jornadas, 5.000 cartillas, 60 custodios certificados y convenios de custodia", "1.572.190.872"),
    ], [5.2, 9.0, 2.8])
    p(d, "Alcance: dos municipios (Manizales, Caldas, y Calarcá, Quindío), cuatro consejos de gestión del riesgo (dos municipales y dos departamentales), "
         "despliegue en tres tramos (30, 40 y 30 nodos) con compuertas de decisión en los meses 12, 17 y 24. El proyecto mide y alerta; no ejecuta ni "
         "autoriza demoliciones y no sustituye la medición regulatoria de la autoridad ambiental.")
    p(d, "Justificación: tras el sismo, los consejos deben decidir alertas y protecciones en los sitios de remoción con sensores de bajo costo cuyo error "
         "no se conoce. La humedad y el polvo grueso alteran esas lecturas. Una alerta con dato de error desconocido puede ser tardía o falsa. El proyecto "
         "cierra esa brecha de ciencia, tecnología e innovación con una red verificada y un uso institucional documentado.")

    h(d, "3. Productos y resultados esperados")
    tabla(d, ["Productos (catálogo MGA, programa 3906)", "Resultados verificables al mes 24"], [
        ("Dos prototipos validados: sistema de 100 nodos y 10 gateways, y gemelo predictivo de demolición", "Error de los nodos: R² ≥ 0,70 y NRMSE ≤ 30 % frente a la referencia"),
        ("Infraestructura de medición y laboratorio de modelación", "Disponibilidad mensual de la red ≥ 90 %"),
        ("Tres documentos de planeación: protocolo de alerta, plan de sostenibilidad y protocolo de evaluación previa de demolición", "Al menos 3 de los 4 consejos deciden con el tablero"),
        ("29 actividades de apropiación social y 5.000 cartillas", "Reducción de falsas alarmas de 30 % o más"),
        ("Asistencia técnica: 60 personas certificadas y 4 consejos", "60 custodios certificados y 12 evaluaciones previas de demolición (al menos 3 por municipio)"),
        ("Nueve documentos de investigación; formación de 6 jóvenes investigadores, 2 magísteres y 1 posdoctor", "A 12 meses del cierre: 70 % de la red operando bajo custodia"),
    ], [8.5, 8.5])

    h(d, "4. Entidad proponente y entidades aliadas")
    tabla(d, ["Entidad", "Calidad", "Función principal", "Contrapartida (COP)"], [
        ("Universidad Santo Tomás, sede Villavicencio", "Proponente y ejecutora", "Administra los recursos del SGR; lidera la medición, la red, la plataforma y el laboratorio", "1.244.000.000 en especie"),
        ("Municipio de Manizales", "Aliada No. 1 (entidad territorial)", "Uso institucional del tablero y del protocolo en su jurisdicción; acceso a sitios; custodia local", "60.000.000 en especie (software y hardware)"),
        ("Municipio de Calarcá", "Aliada No. 2 (entidad territorial)", "Las mismas funciones en su jurisdicción", "60.000.000 en especie (software y hardware)"),
        ("Avanciencia", "Aliada No. 3 (actor del SNCTI)", "Apropiación social, formación de custodios, convenios y propiedad intelectual, por contrato con la universidad", "150.000.000 en especie"),
        ("Gestión y Medio Ambiente", "Aliada No. 4 (sector productivo)", "Equipos y apoyo técnico en Manizales; logística de instalación", "110.000.000 en especie y 150.000.000 en efectivo"),
    ], [4.3, 3.2, 6.3, 3.2])

    h(d, f"5. Qué se compromete a hacer el Municipio de {mun}")
    for x in (f"Liderar en su jurisdicción el uso institucional del tablero y del protocolo de alerta, a través del consejo municipal de gestión del riesgo, y designar delegados para 20 talleres de co-diseño y la formación (actividades 3.1 a 3.5).",
              "Facilitar el acceso a los sitios de remoción y a las demoliciones, e informar las demoliciones previstas para validar el modelo predictivo.",
              f"Aportar en especie 60.000.000 COP (partida {eq} del Anexo 3 v7) mediante el uso de software y hardware existentes durante la ejecución: licencias del sistema de información geográfica para dos puestos (20.000.000), alojamiento del tablero y de los datos en el centro de datos municipal (15.000.000), equipos de la sala de gestión del riesgo (15.000.000) y conectividad y energía en las sedes donde se instalen nodos (10.000.000). No incluye horas de funcionarios ni desembolsos en efectivo.",
              "Adoptar el protocolo de evaluación previa de demolición como insumo de decisión (no como autorización).",
              "Firmar al mes 17 un convenio de custodia de los nodos de su jurisdicción, con personal designado y una partida presupuestal para su operación después del cierre."):
        bullet(d, x)
    p(d, "El municipio no administra recursos del SGR, no figura como beneficiario y no asume obligaciones de compra. La universidad compra los equipos y los "
         "entrega por acta al custodio.", s=9.5, i=True)

    h(d, "6. Valor total, fuentes de financiación y contrapartidas")
    tabla(d, ["Fuente", "Valor (COP)", "Participación"], [
        ("Sistema General de Regalías, Asignación para la CTeI (solicitado)", "13.999.405.215", "88,75 %"),
        ("Contrapartida en especie (universidad, municipios, Avanciencia, Gestión y Medio Ambiente)", "1.624.000.000", "10,30 %"),
        ("Contrapartida en efectivo (Gestión y Medio Ambiente)", "150.000.000", "0,95 %"),
        ("Total del proyecto", "15.773.405.215", "100 %"),
    ], [10.5, 3.5, 3.0])
    p(d, f"Inversión localizada: el Anexo 3 localiza de forma explícita 399.199.591 COP del SGR en cada municipio (técnicos de campo, estación meteorológica, "
         f"estudio de cobertura y talleres); el resto (94,3 %) es común a los dos municipios e incluye los 50 nodos y 5 gateways que se instalan en {mun}.")
    p(d, "Costo después del cierre: el SGR no financia la operación posterior. El costo anual calculable de toda la red es de 265.716.408 COP (licencias, "
         "nube, conectividad satelital y repuestos) y se reparte entre la universidad y los municipios según los convenios de custodia. La parte del municipio "
         "se fija en su convenio al mes 17 y debe contar con partida presupuestal.", b=False)

    h(d, "7. Aliados estratégicos y actores relacionados")
    for x in ("Consejos municipales y departamentales de gestión del riesgo de Caldas y del Quindío: usuarios del tablero y del protocolo.",
              "Gobernaciones de Caldas y del Quindío: certificado de concordancia con el plan de acción específico (si lo suscribieron) y compromiso de uso de sus consejos.",
              "Autoridades ambientales (Corpocaldas y CRQ) y UNGRD: reciben el reporte técnico de desempeño de los sensores; no son integrantes.",
              "Laboratorio acreditado y evaluador externo independiente: verifican el desempeño (TRL 8 al mes 23)."):
        bullet(d, x)

    h(d, "8. Beneficios para el municipio")
    for x in (f"Mediciones de PM2,5 y PM10 con error conocido en los sitios de remoción de {mun} y un tablero con índice local de calidad del aire.",
              "Protocolo de alerta co-diseñado con el consejo municipal y personal formado como custodio.",
              "Evaluaciones previas de la exposición al polvo antes de las demoliciones de su jurisdicción.",
              "Participación en el comité directivo de la alianza y en las decisiones de cada compuerta."):
        bullet(d, x)

    h(d, "9. Riesgos que el municipio debe conocer")
    tabla(d, ["Riesgo", "Cómo se maneja"], [
        ("Que el consejo no adopte el tablero (probabilidad estimada 40 %)", "Co-diseño de umbrales y uso piloto; la compuerta del mes 12 exige que al menos la mitad de los consejos del tramo 1 lo use antes de ampliar"),
        ("Que el custodio no asuma la operación al cierre (50 %)", "Convenio con partida exigido al mes 17; custodia alterna; cláusula de apertura de datos y código"),
        ("Robo o vandalismo de nodos (40 %)", "Cerramientos, custodia local por sitio y 20 nodos de reserva"),
    ], [7.0, 10.0])

    h(d, "10. Documentos que se solicita firmar al señor Alcalde")
    tabla(d, ["Documento", "Para qué", "Fecha límite"], [
        ("Anexo 1: carta unificada de aval, compromiso y gobernanza", "Requisito de la convocatoria (numeral 10.4); la firman los cinco integrantes", "Domingo 1 de noviembre"),
        ("Certificado de participación (numeral 10.18)", "Acredita las actividades y la capacidad técnica del municipio", "Domingo 1 de noviembre"),
        (f"Certificación de contrapartida en especie ({eq})", "Acredita los 60.000.000 COP en software y hardware", "Domingo 1 de noviembre"),
        ("Certificado de precios unitarios", "Respalda los costos del presupuesto en el territorio", "Domingo 1 de noviembre"),
        ("Compromiso de uso (numeral 10.6)", "Respalda el uso del tablero por el consejo municipal", "Domingo 1 de noviembre"),
        ("Certificado de concordancia con el plan de acción específico (numeral 10.21), si el municipio suscribió el plan", "Requisito no subsanable por lugar de ejecución", "Domingo 1 de noviembre"),
        ("Plan de sostenibilidad (numeral 10.16)", "Compromiso de operación posterior del tablero y de los nodos", "Domingo 1 de noviembre"),
    ], [7.0, 7.0, 3.0])
    p(d, "Firmar estos documentos no compromete recursos distintos de la contrapartida en especie descrita y de la partida de custodia que se fije en el convenio del mes 17.", s=9.5, i=True)
    p(d, "Contacto del proyecto: [Diligenciar: nombre, cargo, correo y teléfono del responsable de la Universidad Santo Tomás].", s=9.5)

    if os.path.exists(GRAF):
        sec = d.add_section(); sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = Cm(1.2); sec.top_margin = sec.bottom_margin = Cm(1.0)
        p(d, "Anexo: esquema integrado del proyecto (red del problema a la solución, árbol de decisión, marco conceptual y compromisos de cada actor)", b=True, s=9, color=AZUL)
        d.add_picture(GRAF, height=Cm(17.0))
    f = os.path.join(OUT, f"Resumen_ejecutivo_Alcaldia_de_{mun.replace('á', 'a')}.docx")
    os.makedirs(OUT, exist_ok=True); d.save(f); print("ok", f)

if __name__ == "__main__":
    resumen("Manizales", "Caldas", "EQ-24", "Calarcá")
    resumen("Calarcá", "Quindío", "EQ-25", "Manizales")
