# -*- coding: utf-8 -*-
"""Modelos de certificados para radicar (numerales 10.6, 10.10, 10.18, 10.19 y 10.21 de los TdR de la Convocatoria 56).
Uso: python3 build/modelos_radicacion.py carpeta_salida"""
import os, sys
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = sys.argv[1]
TIT = ("Implementación de una red IoT híbrida Edge-Cloud y gemelo digital para la gestión anticipada del riesgo "
       "por material particulado posterremoto")
PROY = f"«{TIT}» (código SIGP 143329), presentado a la Convocatoria 56 del Sistema General de Regalías, Asignación para la CTeI"

def nuevo():
    d = docx.Document()
    s = d.styles["Normal"]; s.font.name = "Arial"; s.font.size = Pt(10.5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.2)
    return d

def p(d, t, b=False, al=None, s=10.5, after=6):
    q = d.add_paragraph(); r = q.add_run(t); r.bold = b; r.font.size = Pt(s)
    q.paragraph_format.space_after = Pt(after)
    if al is not None: q.alignment = al
    return q

def firma(d, ent, cargo="Representante legal"):
    p(d, "\n\nFirma: ______________________________", after=2)
    p(d, "[Diligenciar: nombre completo]", b=True, after=2)
    p(d, f"{cargo} · C. C. [Diligenciar: número]", after=2)
    p(d, f"{ent} · NIT [Diligenciar]", after=2)

def tabla(d, enc, filas):
    t = d.add_table(rows=1, cols=len(enc)); t.style = "Table Grid"
    for i, x in enumerate(enc):
        c = t.rows[0].cells[i]; c.text = x
        for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(9)
    for f in filas:
        cs = t.add_row().cells
        for i, x in enumerate(f):
            cs[i].text = x
            for r in cs[i].paragraphs[0].runs: r.font.size = Pt(9)
    d.add_paragraph()

def guardar(d, ruta):
    os.makedirs(os.path.dirname(ruta), exist_ok=True); d.save(ruta); print("ok", ruta)

# ------------------------------------------------------------------ 10.18 certificado de participación
INTEG = {
 "Universidad Santo Tomás, sede Villavicencio": ("entidad proponente y ejecutora",
   [("OE1 y OE2 (actividades 1.1 a 1.8 y 2.1 a 2.8)", "Lidera la medición verificada, la red de 100 nodos y 10 gateways, la plataforma y el laboratorio de modelación; ejecuta y administra los recursos del SGR"),
    ("OE3 (actividades 3.1 a 3.6)", "Contrata con Avanciencia la apropiación social y la transferencia; aporta los coinvestigadores de ambiente y ciencias sociales; custodia permanente de los datos")],
   "$ 1.244.000.000 en especie", "Grupo de investigación [Diligenciar: nombre y categoría]; experiencia en [Diligenciar: dos proyectos de CTeI del Anexo 5]; laboratorios y equipo de investigación descritos en la sección 26 del Anexo 2"),
 "Municipio de Manizales": ("entidad territorial aliada No. 1",
   [("OE3 (actividades 3.1 a 3.5)", "Lidera en su jurisdicción el uso institucional del tablero y del protocolo de alerta y adopta el protocolo de evaluación previa de demolición"),
    ("OE1 y OE2 (actividades 1.1, 1.3 a 1.7, 2.1, 2.3 a 2.5)", "Facilita el acceso a los sitios de remoción y a las demoliciones e informa las demoliciones previstas; custodia local de los nodos de su jurisdicción")],
   "$ 60.000.000 en especie", "Consejo municipal de gestión del riesgo en funcionamiento (Ley 1523 de 2012); plan municipal de gestión del riesgo [Diligenciar: acto de adopción]"),
 "Municipio de Calarcá": ("entidad territorial aliada No. 2",
   [("OE3 (actividades 3.1 a 3.5)", "Lidera en su jurisdicción el uso institucional del tablero y del protocolo de alerta y adopta el protocolo de evaluación previa de demolición"),
    ("OE1 y OE2 (actividades 1.1, 1.3 a 1.7, 2.1, 2.3 a 2.5)", "Facilita el acceso a los sitios de remoción y a las demoliciones e informa las demoliciones previstas; custodia local de los nodos de su jurisdicción")],
   "$ 60.000.000 en especie", "Consejo municipal de gestión del riesgo en funcionamiento (Ley 1523 de 2012); plan municipal de gestión del riesgo [Diligenciar: acto de adopción]"),
 "Avanciencia": ("entidad aliada No. 3, actor del SNCTI (oficina de transferencia de resultados de investigación)",
   [("OE1 (actividades 1.1, 1.2, 1.4 y 1.6)", "Vigilancia tecnológica, análisis de libertad de operación, documentación del modelo como resultado protegible y valoración de resultados"),
    ("OE3 (actividades 3.1 a 3.5)", "Ejecuta, por contrato con la universidad, la apropiación social (co-diseño, jornadas, cartillas y publicaciones), forma a 60 custodios, estructura los convenios y gestiona la propiedad intelectual y los registros")],
   "$ 150.000.000 en especie", "Reconocimiento de MinCiencias como actor del SNCTI [Diligenciar: resolución y vigencia]; procesos de transferencia acompañados [Diligenciar]"),
 "Gestión y Medio Ambiente": ("entidad aliada No. 4, actor del sector productivo",
   [("OE1 (actividades 1.1, 1.3 y 1.6)", "Aporta equipos de monitoreo de calidad del aire y apoyo técnico al diagnóstico, la colocalización y la verificación final en Manizales"),
    ("OE2 (actividades 2.3 y 2.4)", "Financia en efectivo la logística y los viáticos de la instalación de los tramos 2 y 3")],
   "$ 260.000.000: $ 110.000.000 en especie y $ 150.000.000 en efectivo", "Empresa de servicios de calidad del aire con domicilio en Manizales, constituida el [Diligenciar: fecha, al menos un año antes del 8 de octubre de 2026]; equipos y experiencia en monitoreo [Diligenciar]"),
}
def cert_1018(ent, carpeta):
    rol, acts, con, cap = INTEG[ent]
    d = nuevo()
    p(d, f"[Diligenciar: membrete de {ent}]", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, "CERTIFICADO DE PARTICIPACIÓN", b=True, al=WD_ALIGN_PARAGRAPH.CENTER, s=12)
    p(d, "Numeral 10.18 de los términos de referencia de la Convocatoria 56", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, f"El suscrito representante legal de {ent} certifica que la entidad participa como {rol} en el proyecto {PROY}, cuya entidad proponente y ejecutora es la Universidad Santo Tomás, sede Villavicencio, y que se compromete a ejecutar las actividades siguientes:")
    tabla(d, ["Objetivo y actividades", "Compromiso"], acts)
    p(d, f"Contrapartida comprometida: {con}, según el Anexo 3 versión 6 y la certificación de contrapartida que se adjunta.")
    p(d, f"Justificación de la capacidad técnica: {cap}.")
    p(d, "Responsable institucional designado: [Diligenciar: nombre, cargo, correo y teléfono].")
    p(d, "Se expide en [Diligenciar: ciudad], el [Diligenciar: día] de noviembre de 2026.")
    firma(d, ent)
    guardar(d, os.path.join(OUT, carpeta, f"Certificado_participacion_10.18_{ent.split(',')[0].replace(' ', '_')}.docx"))

# ------------------------------------------------------------------ 10.19 titularidad
def cert_1019(carpeta):
    d = nuevo()
    p(d, "[Diligenciar: membrete de la Universidad Santo Tomás]", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, "CERTIFICADO DE TITULARIDAD, ADMINISTRACIÓN, CUSTODIA Y CUIDADO DE LOS BIENES", b=True, al=WD_ALIGN_PARAGRAPH.CENTER, s=12)
    p(d, "Numeral 10.19 de los términos de referencia de la Convocatoria 56", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, f"Los suscritos representantes legales de las entidades de la alianza del proyecto {PROY} certifican que concertaron la titularidad, la administración, la custodia y el cuidado de los bienes adquiridos con recursos del SGR y de los resultados, así:")
    tabla(d, ["Componente", "Titular durante la ejecución", "Custodio al cierre (mes 24)", "Condición de recepción"], [
        ("Nodos, gateways y repuestos (120 nodos, 11 gateways)", "Universidad Santo Tomás", "Municipios de Manizales y de Calarcá para los de su jurisdicción; universidad como custodia alterna", "Convenio con partida, personal designado y competencia jurídica"),
        ("Conjunto de referencia, banco de calibración y protocolo de calibración", "Universidad Santo Tomás", "Universidad Santo Tomás; autoridad ambiental regional como candidata adicional", "Personal certificado y plan de mantenimiento"),
        ("Plataforma, gemelo digital, tablero y datos", "Universidad Santo Tomás", "Universidad Santo Tomás (datos y plataforma); municipios como usuarios del tablero", "Licencia de uso y convenio de operación"),
        ("Laboratorio de modelación y gemelo predictivo", "Universidad Santo Tomás", "Universidad Santo Tomás, en la sede de Manizales o Calarcá [Diligenciar]", "Compromiso de sostenimiento del laboratorio"),
        ("Software, modelos y bases de datos (propiedad intelectual)", "Según el acuerdo de titularidad del mes 6", "Según el acuerdo de titularidad; Avanciencia gestiona registros y licencias", "Registro ante la DNDA y licencia de uso a los custodios"),
    ])
    p(d, "Los bienes permanecen en los municipios afectados en los que se ejecuta el proyecto. La proporción de titularidad de los resultados entre los integrantes es [Diligenciar: proporción acordada]. Si un custodio no cumple la condición de recepción, el componente pasa a la custodia alterna y, en su defecto, se aplica la cláusula de apertura del acuerdo marco.")
    p(d, "Se expide en [Diligenciar: ciudad], el [Diligenciar: día] de noviembre de 2026.")
    for e in ("Universidad Santo Tomás, sede Villavicencio", "Municipio de Manizales", "Municipio de Calarcá", "Avanciencia", "Gestión y Medio Ambiente"):
        firma(d, e)
    guardar(d, os.path.join(OUT, carpeta, "Certificado_titularidad_custodia_10.19.docx"))

# ------------------------------------------------------------------ 10.21 concordancia con el plan de acción
def cert_1021(ent, carpeta, alcance):
    d = nuevo()
    p(d, f"[Diligenciar: membrete de {ent}]", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, "CERTIFICADO DE CONCORDANCIA CON EL PLAN DE ACCIÓN ESPECÍFICO", b=True, al=WD_ALIGN_PARAGRAPH.CENTER, s=12)
    p(d, "Numeral 10.21 de los términos de referencia de la Convocatoria 56", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, f"El suscrito representante legal de {ent}, entidad que suscribió el plan de acción específico [Diligenciar: nombre y fecha del plan] para la recuperación tras el sismo del 10 de agosto de 2026, certifica que el proyecto {PROY} es concordante con dicho plan en {alcance}.")
    p(d, "Certifica además que el proyecto no reactiva el riesgo preexistente, porque: (i) instala nodos de medición sobre estructuras existentes o bases de obra menor, sin intervenir cauces, laderas ni edificaciones; (ii) no demuele ni autoriza demoliciones, sino que mide y modela las que ejecutan otros; y (iii) entrega a los consejos de gestión del riesgo información con error conocido para decidir alertas y medidas de protección.")
    p(d, "Línea o acción del plan con la que concuerda: [Diligenciar: línea, programa o acción del plan de acción específico].")
    p(d, "Se expide en [Diligenciar: ciudad], el [Diligenciar: día] de noviembre de 2026.")
    firma(d, ent)
    guardar(d, os.path.join(OUT, carpeta, f"Certificado_concordancia_plan_accion_10.21_{ent.replace(' ', '_').replace('ó','o').replace('í','i').replace('á','a')}.docx"))

# ------------------------------------------------------------------ 10.6 compromiso de uso
def comp_106(ent, carpeta, consejo):
    d = nuevo()
    p(d, f"[Diligenciar: membrete de {ent}]", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, "COMPROMISO DE USO Y ADOPCIÓN DE LOS RESULTADOS", b=True, al=WD_ALIGN_PARAGRAPH.CENTER, s=12)
    p(d, "Numeral 10.6 y Anexo 14 de los términos de referencia de la Convocatoria 56", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, f"El suscrito representante legal de {ent} manifiesta que, en caso de que el proyecto {PROY} resulte financiado, la entidad se compromete a:")
    for x in (f"Designar delegados del {consejo} para participar en el co-diseño de los umbrales y del protocolo de alerta (actividad 3.1) y en el uso del tablero (actividad 3.2).",
              "Registrar las decisiones de alerta y de protección tomadas con el tablero, para verificar el resultado R3 (uso por al menos 75 % de los consejos al mes 24).",
              "Recibir, según el convenio de custodia, los componentes que le correspondan, con partida presupuestal y personal designado para su operación después del cierre.",
              "Solicitar y usar las evaluaciones previas de demolición conforme al protocolo que adopte la actividad 3.5, como insumo de decisión y no como autorización."):
        p(d, "• " + x)
    p(d, "Este compromiso no genera erogaciones adicionales a las que la entidad apruebe en su presupuesto ordinario para la custodia.")
    p(d, "Se expide en [Diligenciar: ciudad], el [Diligenciar: día] de noviembre de 2026.")
    firma(d, ent)
    guardar(d, os.path.join(OUT, carpeta, f"Compromiso_uso_10.6_{ent.replace(' ', '_').replace('ó','o').replace('í','i').replace('á','a')}.docx"))

# ------------------------------------------------------------------ 10.10 solicitud de aval de ética
def solicitud_etica(carpeta):
    d = nuevo()
    p(d, "[Diligenciar: membrete de la Universidad Santo Tomás]", al=WD_ALIGN_PARAGRAPH.CENTER, s=9)
    p(d, "[Diligenciar: ciudad], [Diligenciar: fecha]")
    p(d, "Señores\nComité de Ética de la Investigación\nUniversidad Santo Tomás\n[Diligenciar: sede]")
    p(d, "Asunto: Solicitud de aval ético para el proyecto SIGP 143329 (Convocatoria 56 SGR; cierre 3 de noviembre de 2026)", b=True)
    p(d, f"Solicitamos el aval del Comité para el proyecto {PROY}. El proyecto no realiza experimentos con personas ni recolecta muestras biológicas, pero trata datos personales de participantes en 20 talleres, nueve jornadas públicas y encuestas de percepción (nombres, contacto, imagen y voz con autorización separada, sexo, edad y, de forma voluntaria, pertenencia étnica). Por esa razón el aval es obligatorio según el numeral 10.10 de los términos de referencia, y su falta al cierre es causal de rechazo (numeral 11.6).")
    p(d, "Adjuntamos: (1) protocolo de participación y de co-diseño; (2) formato de consentimiento informado y autorización de tratamiento de datos (Ley 1581 de 2012); (3) tabla de datos personales tratados y medidas de minimización (sección 25 del Anexo 2); (4) protocolo de seguridad del personal de campo.")
    p(d, "Solicitamos que el aval se acompañe del acto administrativo de conformación del Comité, que también se carga con la postulación.")
    p(d, "Cordialmente,")
    p(d, "\n\n[Diligenciar: investigador principal]", b=True, after=2)
    p(d, "Investigador principal · Universidad Santo Tomás", after=2)
    guardar(d, os.path.join(OUT, carpeta, "Solicitud_aval_comite_etica_10.10.docx"))

if __name__ == "__main__":
    U, M, C, A, G, GOB = ("02_FIRMAR/A_Universidad_Santo_Tomas", "02_FIRMAR/B_Municipio_de_Manizales", "02_FIRMAR/C_Municipio_de_Calarca",
                          "02_FIRMAR/D_Avanciencia", "02_FIRMAR/E_Gestion_y_Medio_Ambiente", "02_FIRMAR/F_Gobernaciones_de_Caldas_y_Quindio")
    for ent, c in (("Universidad Santo Tomás, sede Villavicencio", U), ("Municipio de Manizales", M), ("Municipio de Calarcá", C), ("Avanciencia", A), ("Gestión y Medio Ambiente", G)):
        cert_1018(ent, c)
    cert_1019("02_FIRMAR/00_Firma_conjunta_cinco_integrantes")
    cert_1021("Municipio de Manizales", M, "lo que corresponde al municipio de Manizales")
    cert_1021("Municipio de Calarcá", C, "lo que corresponde al municipio de Calarcá")
    cert_1021("Gobernación de Caldas", GOB, "lo que corresponde a los municipios del departamento de Caldas incluidos de forma expresa en el plan")
    cert_1021("Gobernación del Quindío", GOB, "lo que corresponde a los municipios del departamento del Quindío incluidos de forma expresa en el plan")
    comp_106("Municipio de Manizales", M, "consejo municipal de gestión del riesgo de Manizales")
    comp_106("Municipio de Calarcá", C, "consejo municipal de gestión del riesgo de Calarcá")
    comp_106("Gobernación de Caldas", GOB, "consejo departamental de gestión del riesgo de Caldas")
    comp_106("Gobernación del Quindío", GOB, "consejo departamental de gestión del riesgo del Quindío")
    solicitud_etica(U)
