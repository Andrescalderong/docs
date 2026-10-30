# -*- coding: utf-8 -*-
"""Diligencia el Anexo 1 oficial (carta unificada de aval, compromiso institucional y modelo de gobernanza) con el Anexo 3 v6.
Uso: python3 build/anexo1_v6.py plantilla.docx salida.docx"""
import sys, re, copy, docx
from docx.shared import RGBColor
from docx.oxml.ns import qn

src, dst = sys.argv[1:3]
d = docx.Document(src)
P = d.paragraphs
NEGRO, VERDE = RGBColor(0, 0, 0), RGBColor(0, 0xCC, 0)
TIT = ("Implementación de una red IoT híbrida Edge-Cloud y gemelo digital para la gestión anticipada del riesgo "
       "por material particulado posterremoto")
CONV = ("Convocatoria de emergencia para la recuperación y generación de capacidades transformadoras, en el marco "
        "del Decreto 1414 de 17 de septiembre de 2026 (Convocatoria 56 del SGR)")

def poner(p, txt):
    """Reemplaza el texto del párrafo conservando el formato de su primer run; las marcas [Diligenciar] quedan en verde."""
    runs = p.runs
    base = copy.deepcopy(runs[0]._r) if runs else None
    for r in runs: r._r.getparent().remove(r._r)
    for frag in re.split(r"(\[Diligenciar:[^\]]*\])", txt):
        if not frag: continue
        r = p.add_run(frag)
        if base is not None and base.rPr is not None:
            r._r.insert(0, copy.deepcopy(base.rPr))
        else:
            from docx.shared import Pt
            r.font.name = "Arial"; r.font.size = Pt(10)
        r.font.color.rgb = VERDE if frag.startswith("[Diligenciar") else NEGRO
        r.font.highlight_color = None

def borrar(p):
    p._p.getparent().remove(p._p)

# ------------------------------------------------------------------ carta
poner(P[11], "[Diligenciar: ciudad], [Diligenciar: día] de noviembre de 2026")
poner(P[19], f"Asunto: Aval, compromiso institucional y modelo de gobernanza del proyecto «{TIT}», código SIGP 143329.")
poner(P[24], "Por medio de la presente yo [Diligenciar: nombre del representante legal], identificado(a) con cédula de ciudadanía número "
      "[Diligenciar: número], como representante legal de la Universidad Santo Tomás, sede Villavicencio, avalo el proyecto titulado "
      f"«{TIT}», a la {CONV}, el cual cumple con criterios de calidad científico-técnica. Manifiesto que las actividades para las cuales "
      "se solicita la financiación de recursos mediante la Asignación para la Ciencia, Tecnología e Innovación del Sistema General de "
      "Regalías no están siendo actualmente financiadas con recursos del Ministerio de Ciencia, Tecnología e Innovación, del Sistema "
      "General de Regalías (SGR), o con otras fuentes de financiación. Como también aclaro que, en el marco de la presente convocatoria, "
      "no participo en el rol de proponente en más de un (1) proyecto.")
poner(P[28], "Para el desarrollo del proyecto se conformará la Alianza [Diligenciar: nombre de la alianza; propuesta: Alianza Aire Eje Cafetero], "
      "que estará integrada por las siguientes entidades: Universidad Santo Tomás, sede Villavicencio; Municipio de Manizales (Caldas); "
      "Municipio de Calarcá (Quindío); Avanciencia [Diligenciar: razón social]; y Gestión y Medio Ambiente [Diligenciar: razón social]; "
      "designándose como entidad proponente a la Universidad Santo Tomás, quien será la encargada de la ejecución técnica y financiera "
      "ante el Órgano Colegiado de Administración y Decisión (OCAD) de Ciencia, Tecnología e Innovación (CTeI) del Sistema General de "
      "Regalías (SGR), en caso de que el proyecto de inversión sea viabilizado, priorizado y aprobado por el OCAD de CTeI del SGR.")
poner(P[34], "Informo a MINCIENCIAS que las personas encargadas de ejercer la función de interlocutores válidos del proyecto son:")
poner(P[36], "Nombre del interlocutor: [Diligenciar: nombre de la persona que administra los recursos]")
poner(P[37], "Documento de identidad: [Diligenciar: tipo y número]")
poner(P[38], "Correo electrónico: [Diligenciar: correo]")
poner(P[39], "Teléfono: [Diligenciar: teléfono]")
poner(P[41], "Nombre del líder o investigador principal: [Diligenciar: nombre del investigador principal]")
poner(P[42], "Documento de identidad: [Diligenciar: tipo y número]")
poner(P[43], "Correo electrónico: [Diligenciar: correo]")
poner(P[44], "Teléfono: [Diligenciar: teléfono]")
poner(P[48], "Aspectos éticos:")
poner(P[53], "Informo al Ministerio que el proyecto presentado a la Convocatoria 56 del SGR de 2026 fue revisado por el Comité de Ética de "
      "la Universidad Santo Tomás, debidamente constituido. Se adjunta el aval del Comité de Ética y el acto administrativo de constitución "
      "del mismo [Diligenciar: número y fecha del aval]. El proyecto trata datos personales en talleres, encuestas y registros de asistencia, "
      "con consentimiento informado según la Ley 1581 de 2012.")
poner(P[67], "El tiempo estimado para la ejecución técnica y financiera del proyecto es de: veinticuatro (24) meses.")
poner(P[90], f"Además, ACEPTAMOS expresa e irrevocablemente que conocemos detalladamente las características, requisitos y condiciones de la "
      f"{CONV}, de manera que nos sometemos a lo establecido en los Términos de Referencia y los anexos determinados por MINCIENCIAS para "
      "el desarrollo de la misma.")
poner(P[95], "Las partes abajo firmantes convienen la intención de acuerdo que se regirá bajo las siguientes cláusulas: i) la titularidad de la "
      "propiedad intelectual sobre los resultados que se obtengan o se pudieran obtener en el desarrollo del presente proyecto estará a cargo "
      "de la Universidad Santo Tomás y de las entidades aliadas que aporten conocimiento o recursos al resultado, en la proporción que fije "
      "el acuerdo marco de titularidad del mes 6 del proyecto. ii) La distribución de los derechos patrimoniales sobre todos y cada uno de los "
      "entregables generados se establecerá de acuerdo con el porcentaje de los aportes desembolsables y no desembolsables que las partes "
      "realicen, los cuales se determinarán en el acta de liquidación del contrato. iii) De igual manera cualquiera de las partes podrá "
      "iniciar los mecanismos de protección correspondientes garantizando el respeto de los porcentajes que a cada una de las partes "
      "corresponden sobre los resultados; por su parte, las demás entidades se comprometen a disponer de los recursos que sean necesarios "
      "para facilitar y apoyar que se conceda por la autoridad competente el mecanismo de protección que se determine. Avanciencia "
      "acompaña la valoración, la protección y el registro de los resultados. iv) La custodia y cuidado de los productos tecnológicos y "
      "prototipos que se materialicen estará a cargo de la Universidad Santo Tomás durante la duración del proyecto. A la finalización y "
      "liquidación del mismo se definirá la custodia y cuidado de tales bienes, con los municipios de Manizales y de Calarcá como custodios "
      "locales de los componentes de su jurisdicción. v) Los derechos morales de autor que correspondan a estudiantes, profesores o "
      "investigadores de las partes, que por sus aportes significativos en una determinada obra les corresponden como autores o coautores, "
      "serán a estos siempre reconocidos. vi) Ninguna de las partes podrá publicar, comunicar, divulgar, revelar ni permitir que los "
      "investigadores y personal vinculado al contrato publiquen, comuniquen, revelen o utilicen la información resultado del mismo, sin "
      "previo aviso y aprobación por escrito. vii) Sin perjuicio de lo anterior, las partes podrán efectuar modificaciones al presente "
      "documento de acuerdo con las condiciones de desarrollo del proyecto.")
poner(P[102], "Las entidades de la alianza se acogen a la Política Nacional de Ética de la Investigación, Bioética e Integridad Científica y a "
      "las políticas de integridad de cada entidad. Acuerdan cinco compromisos mínimos: consentimiento informado y tratamiento de datos "
      "personales conforme a la Ley 1581 de 2012; reconocimiento de autoría según la contribución; declaración y gestión de conflictos de "
      "interés; publicación de los datos de medición con sus métricas de error, sin datos personales; y advertencia expresa de que el "
      "resultado del laboratorio de modelación no constituye una autorización ni una prueba de responsabilidad.")

# firmas: proponente y aliados 1 y 2 del formato; se añaden los aliados 3 y 4
FIRMAS = [
    (113, ["__________________________", "FIRMA", "[Diligenciar: nombre del representante legal]", "C.C. [Diligenciar: número]", "Representante legal",
           "UNIVERSIDAD SANTO TOMÁS, SEDE VILLAVICENCIO (ENTIDAD PROPONENTE)", "[Diligenciar: dirección]", "[Diligenciar: teléfono]"]),
    (124, ["__________________________", "FIRMA", "[Diligenciar: nombre del alcalde]", "C.C. [Diligenciar: número]", "Alcalde",
           "MUNICIPIO DE MANIZALES (ENTIDAD ALIADA 1)", "[Diligenciar: dirección]", "[Diligenciar: teléfono]"]),
    (133, ["__________________________", "FIRMA", "[Diligenciar: nombre del alcalde]", "C.C. [Diligenciar: número]", "Alcalde",
           "MUNICIPIO DE CALARCÁ (ENTIDAD ALIADA 2)", "[Diligenciar: dirección]", "[Diligenciar: teléfono]"]),
]
bloque = [copy.deepcopy(P[i]._p) for i in range(133, 141)]
for i0, textos in FIRMAS:
    for k, t in enumerate(textos):
        poner(P[i0 + k], t)
ancla = P[140]._p
extra = [["__________________________", "FIRMA", "[Diligenciar: nombre del representante legal]", "C.C. [Diligenciar: número]", "Representante legal",
          "AVANCIENCIA [Diligenciar: razón social y NIT] (ENTIDAD ALIADA 3)", "[Diligenciar: dirección]", "[Diligenciar: teléfono]"],
         ["__________________________", "FIRMA", "[Diligenciar: nombre del representante legal]", "C.C. [Diligenciar: número]", "Representante legal",
          "GESTIÓN Y MEDIO AMBIENTE [Diligenciar: razón social y NIT] (ENTIDAD ALIADA 4)", "[Diligenciar: dirección]", "[Diligenciar: teléfono]"]]
for textos in extra:
    vacio = copy.deepcopy(P[141]._p); ancla.addnext(vacio); ancla = vacio
    for k, el in enumerate(bloque):
        nuevo = copy.deepcopy(el); ancla.addnext(nuevo); ancla = nuevo
        poner(docx.text.paragraph.Paragraph(nuevo, P[0]._parent), textos[k])

# párrafos informativos en rojo y opciones que no aplican
for i in sorted([8, 9, 26, 46, 50, 51, 52], reverse=True):
    borrar(P[i])

# ------------------------------------------------------------------ tabla de contrapartida (Anexo 3 v6)
C = [("Universidad Santo Tomás, sede Villavicencio", 1244000000, 0), ("Municipio de Manizales", 60000000, 0),
     ("Municipio de Calarcá", 60000000, 0), ("Avanciencia", 150000000, 0), ("Gestión y Medio Ambiente", 110000000, 150000000)]
TOT = sum(e + f for _, e, f in C)
assert TOT == 1774000000
def cop(n): return "$ " + f"{int(n):,}".replace(",", ".")
def pc(a, b): return (f"{a / b * 100:.2f}".replace(".", ",") + " %") if b else "0,00 %"
t = d.tables[0]
plantilla_fila = copy.deepcopy(t.rows[2]._tr)
for r in list(t.rows)[2:]: r._tr.getparent().remove(r._tr)
filas = [(n, e + f, pc(e + f, TOT), e, pc(e, e + f), f, pc(f, e + f)) for n, e, f in C]
esp, efe = sum(e for _, e, _ in C), sum(f for _, _, f in C)
filas.append(("Total", TOT, "100,00 %", esp, pc(esp, TOT), efe, pc(efe, TOT)))
for fila in filas:
    tr = copy.deepcopy(plantilla_fila); t._tbl.append(tr)
    row = t.rows[-1]
    for j, v in enumerate(fila):
        cel = row.cells[j]
        txt = cop(v) if isinstance(v, int) else v
        poner(cel.paragraphs[0], txt)
        for p in cel.paragraphs[1:]: p._p.getparent().remove(p._p)

from docx.shared import Pt
for row in t.rows[2:]:
    for cel in row.cells:
        for p in cel.paragraphs:
            for r in p.runs: r.font.size = Pt(7.5)

# ------------------------------------------------------------------ modelo de gobernanza
G = d.tables[1]
def celda(i, j, txt):
    c = G.rows[i].cells[j]
    poner(c.paragraphs[0], txt)
    for p in c.paragraphs[1:]: p._p.getparent().remove(p._p)
def fila_ultima(i, txt):
    celda(i, len(G.rows[i].cells) - 1, txt)

fila_ultima(1, TIT)
ENT = [(3, "Universidad Santo Tomás, sede Villavicencio"), (6, "Municipio de Manizales"), (9, "Municipio de Calarcá"),
       (12, "Avanciencia [Diligenciar: razón social]"), (15, "Gestión y Medio Ambiente [Diligenciar: razón social]")]
for i, nom in ENT:
    celda(i, 2, nom)
    r1 = G.rows[i + 1].cells; r2 = G.rows[i + 2].cells
    poner(r1[1].paragraphs[0], "Nit: [Diligenciar: NIT]")
    poner(r1[3].paragraphs[0], "Representante legal: [Diligenciar: nombre]")
    poner(r1[4].paragraphs[0], "E-mail: [Diligenciar: correo]")
    poner(r2[1].paragraphs[0], "Dirección: [Diligenciar: dirección]")
    ciudad = {3: "Villavicencio (Meta)", 6: "Manizales (Caldas)", 9: "Calarcá (Quindío)", 15: "Manizales (Caldas)"}.get(i, "[Diligenciar: ciudad]")
    poner(r2[3].paragraphs[0], f"Ciudad: {ciudad}")
    poner(r2[4].paragraphs[0], "Teléfono: [Diligenciar: teléfono]")
celda(15, 0, "Aliado 4")
fila_ultima(20, "Desarrollar, validar y transferir, en Manizales y en Calarcá, una red de 100 nodos de medición de PM2,5 y PM10 con corrección en el borde, "
            "un gemelo digital y un laboratorio de modelación que estima la exposición esperada antes de cada demolición. La alianza busca que "
            "los consejos de gestión del riesgo decidan alertas con mediciones de error conocido, en el marco de la emergencia del sismo del 10 de agosto de 2026.")
fila_ultima(21, "Universidad Santo Tomás: ejecuta y administra los recursos del SGR; lidera los objetivos 1 y 2, la apropiación social y la custodia "
            "durante la ejecución. Municipios de Manizales y de Calarcá: lideran el uso institucional del tablero y del protocolo de alerta en su "
            "jurisdicción (objetivo 3), facilitan el acceso a sitios y demoliciones y asumen la custodia local. Avanciencia: gestiona la vigilancia "
            "tecnológica, la propiedad intelectual, la valoración y los convenios de transferencia y custodia. Gestión y Medio Ambiente: aporta equipos "
            "de monitoreo y apoyo técnico en el diagnóstico, la colocalización y la verificación final en Manizales, y logística de los tramos 2 y 3. "
            "Cada integrante certifica su contrapartida y designa un responsable institucional.")
fila_ultima(23, "Comité directivo de la alianza, con un delegado con capacidad de decisión de cada integrante [Diligenciar: nombres]. Lo preside la "
            "Universidad Santo Tomás. Decide las compuertas C1 (mes 12), C2 (mes 17) y C3 (mes 24).")
fila_ultima(24, "Gerencia del proyecto, Universidad Santo Tomás [Diligenciar: nombre del gerente].")
fila_ultima(25, "Investigador principal de la Universidad Santo Tomás [Diligenciar: nombre], con un comité técnico integrado por el coinvestigador técnico, "
            "el especialista en metrología, el investigador líder en modelación y delegados técnicos de Avanciencia y de Gestión y Medio Ambiente.")
fila_ultima(26, "Universidad Santo Tomás, como entidad ejecutora, mediante su dependencia financiera [Diligenciar: dependencia y responsable].")
fila_ultima(27, "Mesa de uso institucional con delegados de los municipios de Manizales y de Calarcá, los consejos de gestión del riesgo y la coordinación de "
            "apropiación social. Evaluación externa independiente en los meses 12 y 24.")
fila_ultima(30, "La Universidad Santo Tomás administra el proyecto con su sistema de gestión de proyectos y su manual de contratación. La gerencia lleva el "
            "plan de trabajo, la matriz de riesgos y el registro de permisos, y entrega ocho informes trimestrales a la supervisión. Cada aliado "
            "reporta trimestralmente su contrapartida con soportes. El convenio de la alianza se firma antes del acta de inicio.")
fila_ultima(31, "Los recursos del SGR se ejecutan por la Universidad Santo Tomás en cuenta y centro de costos exclusivos, con conciliación mensual. "
            "Los aliados no reciben recursos del SGR. La contrapartida es de 1.774.000.000 COP: 1.624.000.000 en especie y 150.000.000 en "
            "efectivo de Gestión y Medio Ambiente, según el Anexo 3. La revisoría fiscal y el rubro de seguimiento sustentan el control.")
fila_ultima(34, "La operación posterior se financia con partidas permanentes de las entidades custodias. La Universidad Santo Tomás sostiene la red, "
            "la plataforma, los datos y el laboratorio. Cada municipio sostiene el tablero, el protocolo de alerta y la custodia local de sus nodos. "
            "El costo anual calculable es de 265.716.408 COP (tabla 16.6 del Anexo 2). Cada entidad operadora firma el plan de sostenibilidad "
            "del numeral 10.16. El seguimiento mide la continuidad de 70 % o más de los nodos a los 12 meses del cierre.")
fila_ultima(37, "Las decisiones técnicas se toman en el comité técnico con criterios numéricos fijados antes de medir. El comité directivo aprueba las "
            "compuertas por mayoría calificada de cuatro de cinco integrantes, con dos firmas en cada acta: investigador principal y gerencia.")
fila_ultima(38, "Reunión mensual del comité directivo, reunión quincenal del comité técnico, correo institucional, repositorio documental compartido "
            "y tablero de seguimiento del proyecto.")
fila_ultima(39, "Plan de trabajo con hitos H0 a H11 y compuertas C1 a C3; indicadores de producto y de resultado; informes trimestrales; "
            "evaluación externa en los meses 12 y 24 y ensayo independiente en el mes 23.")
fila_ultima(40, "Arreglo directo en el comité directivo dentro de 15 días hábiles; si persiste, conciliación ante un centro de conciliación; "
            "en último término, la jurisdicción competente. Las controversias no suspenden la ejecución.")
fila_ultima(43, "Convenio de asociación o de cooperación entre los integrantes; contratos de compra y de prestación de servicios celebrados por la "
            "Universidad Santo Tomás conforme a su manual de contratación, con al menos dos cotizaciones por rubro [Diligenciar: modalidad definida por la universidad].")
fila_ultima(44, "Compra para los equipos financiados por el SGR; comodato o convenio de custodia para la entrega a los custodios al cierre; "
            "licencia de uso para el software.")
fila_ultima(45, "Nodos de medición de PM2,5 y PM10 con comunicación de baja potencia y energía solar; gateways con respaldo satelital; conjunto de "
            "monitores de referencia (gravimétrico, beta y óptico); banco de calibración con cámara climática; servidor de cálculo; estaciones "
            "meteorológicas; dron con LiDAR; monitores portátiles; servicios de nube, de plataforma geoespacial, de laboratorio acreditado y de evaluación externa.")
# textos de instrucción en las celdas no obligatorias
for i in (35, 41, 46, 48):
    c = G.rows[i].cells[-1]
    poner(c.paragraphs[0], "No aplica.")
    for p in c.paragraphs[1:]: p._p.getparent().remove(p._p)

# los bloques de firma no se parten entre páginas
dentro = False
for p in d.paragraphs:
    if p.text.startswith("Atentamente"): dentro = True
    if dentro and p.text.strip() and not p.text.startswith("[Diligenciar: teléfono]"):
        p.paragraph_format.keep_with_next = True
d.save(dst); print("ok", dst)
