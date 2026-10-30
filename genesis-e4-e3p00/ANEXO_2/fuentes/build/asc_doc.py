# -*- coding: utf-8 -*-
"""Extrae la sección 19 del Anexo 2 compilado como documento independiente de la estrategia ASC, con ficha y datos requeridos.
Uso: python3 build/asc_doc.py anexo2.docx salida.docx"""
import sys, copy, docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
src, dst = sys.argv[1:3]
d = docx.Document(src)
body = d.element.body
els = list(body)
def texto(e): return "".join(t.text or "" for t in e.iter(qn("w:t")))
def es_h1(e):
    if e.tag != qn("w:p"): return False
    ps = e.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps is not None and ps.get(qn("w:val")) in ("Heading1", "Ttulo1")
i0 = next(i for i, e in enumerate(els) if es_h1(e) and texto(e).startswith("19."))
i1 = next(i for i, e in enumerate(els) if i > i0 and es_h1(e))
for k, e in enumerate(els):
    if e.tag == qn("w:sectPr"): continue
    if not (i0 <= k < i1): body.remove(e)
for e in list(body):
    s = e.find(qn("w:pPr") + "/" + qn("w:sectPr")) if e.tag == qn("w:p") else None
    if s is not None: s.getparent().remove(s)
VERDE = RGBColor(0, 0x80, 0)
primero = body[0]
ESTILO_T = d.tables[0].style
def p_antes(txt, bold=False, size=10, italic=False):
    p = d.add_paragraph()
    import re
    for frag in re.split(r"(\[Diligenciar:[^\]]*\])", txt):
        if not frag: continue
        r = p.add_run(frag); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = "Arial"
        if frag.startswith("[Diligenciar"): r.font.color.rgb = VERDE
    primero.addprevious(p._p); return p
def t_antes(enc, filas):
    t = d.add_table(rows=1, cols=len(enc)); t.style = ESTILO_T
    for j, h in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5)
    import re
    for f in filas:
        cs = t.add_row().cells
        for j, v in enumerate(f):
            cs[j].text = ""
            for frag in re.split(r"(\[Diligenciar:[^\]]*\])", v):
                if not frag: continue
                r = cs[j].paragraphs[0].add_run(frag); r.font.size = Pt(8.5)
                if frag.startswith("[Diligenciar"): r.font.color.rgb = VERDE
    from docx.oxml import parse_xml
    tblPr = t._tbl.tblPr
    tblPr.append(parse_xml('<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">' + ''.join(f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="808080"/>' for b in ("top", "left", "bottom", "right", "insideH", "insideV")) + '</w:tblBorders>'))
    primero.addprevious(t._tbl)
p_antes("ESTRATEGIA DE APROPIACIÓN SOCIAL DEL CONOCIMIENTO (ASC)", True, 12)
p_antes("Proyecto SIGP 143329. Documento extraído del Anexo 2 (sección 19), con la ficha de la estrategia y los datos que la alianza debe aportar. Sigue el numeral 6.1 de los TdR y la plantilla del Anexo 2; se ajusta al formato del Anexo 10 cuando se disponga de él.", italic=True, size=9)
t_antes(["Campo", "Contenido"], [
 ["Proyecto", "Implementación de una red IoT híbrida Edge-Cloud y gemelo digital para la gestión anticipada del riesgo por material particulado posterremoto"],
 ["Entidad ejecutora", "Universidad Santo Tomás, sede Villavicencio [Diligenciar: NIT]"],
 ["Aliados en la estrategia", "Municipio de Manizales, Municipio de Calarcá y Avanciencia; Gestión y Medio Ambiente como actor técnico externo"],
 ["Territorios", "Manizales (Caldas) y Calarcá (Quindío)"],
 ["Población", "Cuatro consejos de gestión del riesgo, cerca de 40 personas del servicio público, personal operario de remoción y comunidad vecina de los sitios"],
 ["Producto MGA", "Servicio de apropiación social del conocimiento (sector 39, programa 3906): 29 actividades, 20 talleres de co-diseño y 9 jornadas públicas"],
 ["Compromisos asociados", "5.000 cartillas; al menos tres actas firmadas de co-diseño; 60 personas certificadas en uso y custodia (asistencia técnica)"],
 ["Resultado asociado", "Uso del tablero en decisiones documentadas por 75 % o más de los consejos al mes 24"],
 ["Presupuesto", "Actividades 3.1 a 3.4: 1.504.322.462 COP del SGR y 320.000.000 COP de contrapartida en especie"],
 ["Responsables", "Coordinación de apropiación social (Universidad Santo Tomás), con los municipios de Manizales y de Calarcá"],
 ["Periodo", "Meses 3 a 24"],
])
p_antes("Datos e insumos que la alianza debe aportar", True, 10)
t_antes(["Dato o insumo", "Quién lo aporta", "Para qué se usa"], [
 ["Anexo 10 de los TdR (lineamientos de ASC)", "Universidad Santo Tomás", "Verificar el formato y el contenido exigido"],
 ["Barrios y centros poblados con sitios de remoción", "Municipios de Manizales y de Calarcá", "Parte A: reconocimiento del contexto"],
 ["Organizaciones comunitarias y juntas de acción comunal de los barrios", "Municipios", "Parte B: actores y convocatoria"],
 ["Contratistas de remoción y número de operarios por sitio", "Municipios", "Partes B y D: representación del personal operario"],
 ["Certificación de presencia de comunidades étnicas", "Universidad, ante el Ministerio del Interior", "Consulta previa, si aplica"],
 ["Composición de cada consejo de gestión del riesgo", "Municipios y gobernaciones", "Línea base y convocatoria"],
 ["Sedes accesibles para talleres y jornadas", "Municipios", "Logística y ajustes razonables"],
 ["Medios municipales de difusión (emisoras, redes, carteleras)", "Municipios", "Nivel 1 de la escalera"],
 ["Nombres de la coordinación y de los coinvestigadores de ambiente y ciencias sociales", "Universidad Santo Tomás", "Responsables y contrapartida TH-19"],
 ["Responsable del tratamiento de datos personales y política de datos", "Universidad Santo Tomás", "Consentimiento informado (Ley 1581 de 2012)"],
 ["Aval o determinación del comité de ética", "Universidad Santo Tomás", "Talleres y encuestas con personas"],
 ["Experiencias previas de apropiación social de la universidad", "Universidad Santo Tomás", "Idoneidad; las actividades de ASC cuentan como CTeI según la Secretaría Técnica"],
])
d.save(dst); print("ok", dst)
