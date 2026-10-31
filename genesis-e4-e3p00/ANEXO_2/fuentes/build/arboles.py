# -*- coding: utf-8 -*-
"""Formato de plantilla para las tablas de árbol (problemas y objetivos): combina la fila central y colorea por nivel.
Uso: python3 build/arboles.py entrada.docx salida.docx"""
import sys, docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

VERDE, ROSA, DURAZNO = "E2EFDA", "F4CCCC", "FCE4D6"
NIVEL = {"efectos indirectos": VERDE, "efectos directos": VERDE, "fines indirectos": VERDE, "fines directos": VERDE,
         "problema central": ROSA, "objetivo general": ROSA,
         "causas directas": DURAZNO, "causas indirectas": DURAZNO, "medios directos": DURAZNO, "medios indirectos": DURAZNO,
         "objetivos específicos": DURAZNO, "medios": DURAZNO, "causas": DURAZNO}

def sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn("w:shd")): tcPr.remove(s)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def main(src, dst):
    d = docx.Document(src); n = 0
    for t in d.tables:
        c0 = t.rows[0].cells[0].text.strip().lower()
        if c0 not in ("efectos indirectos", "fines indirectos"): continue
        n += 1
        for r in t.rows:
            et = r.cells[0].text.strip().lower()
            col = NIVEL.get(et)
            if not col: continue
            if et in ("problema central", "objetivo general") and len(r.cells) >= 4:
                a, b = r.cells[1], r.cells[len(r.cells) - 1]
                textos = [c.text.strip() for c in r.cells[1:]]
                if all(not x for x in textos[1:]):
                    m = a.merge(b)
                    for p in m.paragraphs[1:]:
                        if not p.text.strip(): p._p.getparent().remove(p._p)
            for c in r.cells: sombra(c, col)
    d.save(dst); print("árboles formateados:", n)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
