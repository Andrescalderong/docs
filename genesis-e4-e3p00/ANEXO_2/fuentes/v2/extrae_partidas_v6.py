import openpyxl, json, re, sys
wb=openpyxl.load_workbook(sys.argv[1],data_only=True)
fc1=json.load(open(sys.argv[2]))
HOJAS={"R01":("01. Talento Humano",9),"R02":("02. Equipos y Software",10),"R03":("03. Capacitación Y Eventos",7),"R04":("04. Servicios Tecnologicos",12),
"R05":("05. Materiales, Insumos y Doc.",10),"R06":("06.Protección conocimiento y Di",9),"R07":("07. Infraestructura",10),"R08":("08. Gastos de viaje",9),
"R09":("09. Administrativos ",7),"R10":("10. Seguimiento",8),"R11":("11. Otros",9)}
ENT=["USTA","Manizales","Calarcá","Avanciencia","GyMA"]
out=[]
for rub,(name,h) in HOJAS.items():
    ws=wb[name]
    hdr={}
    for c in ws[h]:
        if c.value not in (None,""): hdr[c.column_letter]=str(c.value).strip()
    # columnas de fuentes: en fila h (o h+3 para R03) pares Especie/Efectivo por entidad y ultimo Efectivo = SGR
    hrow = h if rub!="R03" else 10
    fuentes=[c.column_letter for c in ws[hrow] if str(c.value).strip() in ("Especie","Efectivo","EFECTIVO")]
    first=hrow+1
    rows=[]
    for r in range(first, ws.max_row+1):
        b=ws[f"B{r}"].value
        if isinstance(b,str) and b.strip().upper().startswith("TOTAL"): break
        vals={hdr.get(c.column_letter,c.column_letter):c.value for c in ws[r] if c.value not in (None,"")}
        if not b and not vals: continue
        src={}
        for i,col in enumerate(fuentes[:-1]):
            v=ws[f"{col}{r}"].value or 0
            if v: src[f"{ENT[i//2]} {'especie' if i%2==0 else 'efectivo'}"]=v
        sgr=ws[f"{fuentes[-1]}{r}"].value or 0
        rows.append(dict(usada=False, fila=r, campos={k:v for k,v in vals.items()}, fuentes_contra=src, sgr=sgr))
    items=[i for i in fc1 if i["rubro"]==rub and not i["id"].startswith("ADIC")]
    rows=[x for x in rows if any(isinstance(v,(int,float)) and v for v in x["campos"].values())]
    for it in items:
        cand=[x for x in rows if str(x["campos"].get(list(hdr.values())[0] if rub!="R01" else "CARGO",""))]
        m=None
        for L in (len(it["item"]),80,60,45,35):
            for x in rows:
                if x.get("usada"): continue
                txt=" ".join(str(v) for v in x["campos"].values())
                if it["item"][:L] in txt: m=x; break
            if m: break
        if m: m["usada"]=True
        if m is None:
            print("SIN FILA", it["id"], it["item"][:60], file=sys.stderr)
        out.append(dict(it, hoja=name, fila=m["fila"] if m else None, campos=m["campos"] if m else {}, fuentes_contra=m["fuentes_contra"] if m else {}))
json.dump(out,open(sys.argv[3],"w"),ensure_ascii=False,indent=1,default=str)
print(len(out))
